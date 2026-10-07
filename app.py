import os
import json
import tempfile
import streamlit as st
import gspread
import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
import pandas as pd
import re
import io
import warnings

# Mengabaikan warning kecil dari Pandas
warnings.filterwarnings('ignore')

# Konfigurasi Halaman Web Streamlit
st.set_page_config(
    page_title="Rekap Absensi Vendor",
    page_icon="📊",
    layout="centered"
)

# Judul Utama
st.title("📊 Aplikasi Otomatisasi Absensi")
st.markdown("Rekap Absensi Vendor Warehouse")
st.divider()

# Fungsi Bantuan
def extract_spreadsheet_id(input_text):
    """Mengekstrak Spreadsheet ID dari link URL"""
    match = re.search(r"/d/([a-zA-Z0-9-_]+)", input_text)
    if match:
        return match.group(1)
    return input_text.strip()

def parse_shift(val):
    """Mengekstrak format Shift"""
    if pd.isna(val):
        return "S1"
    val_str = str(val).strip()
    match = re.search(r'(\d+)', val_str)
    if match:
        return f"S{match.group(1)}"
    return "S1"

# --- UI FORM INPUT ---
link_input = st.text_input("Masukkan Link Google Sheets disini", placeholder="https://docs.google.com/spreadsheets/d/...")
nama_input = st.text_input("Nama File Karyawan", placeholder="Cth: John Doe")

col1, col2 = st.columns(2)
with col1:
    tahun_input = st.selectbox("Pilih Tahun", ["2025", "2026", "2027", "2028", "2029", "2030"], index=1)
with col2:
    bulan_input = st.selectbox("Pilih Bulan", ["01", "02", "03", "04", "05", "06", "07", "08", "09", "10", "11", "12"], index=6)

st.divider()

# --- LOGIKA TOMBOL GENERATE ---
if st.button("🚀 Tarik Data & Generate Excel", use_container_width=True):
    if not link_input:
        st.warning("⚠️ Silakan masukkan link Google Sheets terlebih dahulu!")
    elif not nama_input:
        st.warning("⚠️ Ketik dulu nama karyawannya untuk nama file Excel!")
    else:
        # Menampilkan animasi loading saat memproses
        with st.spinner("Menghubungi Google Server dan Memproses Data..."):
            spreadsheet_id = extract_spreadsheet_id(link_input)
            
            try:
                # 1. TARIK DATA DARI GOOGLE SHEETS
                from creds_helper import get_gspread_client
                gc = get_gspread_client()
                sheet = gc.open_by_key(spreadsheet_id).sheet1
                data_mentah = sheet.get_all_records()
                df_data = pd.DataFrame(data_mentah)
                
                # 2. PROSES PANDAS
                df = df_data.copy()
                df['TANGGAL_ASLI'] = pd.to_datetime(df['TANGGAL'], format='mixed', dayfirst=True, errors='coerce')
                
                # Filter berdasarkan Tahun dan Bulan
                df_filtered = df[
                    (df['TANGGAL_ASLI'].dt.strftime('%Y') == tahun_input) & 
                    (df['TANGGAL_ASLI'].dt.strftime('%m') == bulan_input)
                ].copy()

                if df_filtered.empty:
                    st.error(f"❌ Tidak ada data absen untuk bulan {bulan_input} tahun {tahun_input}")
                else:
                    kolom_timestamp = df_filtered.columns[0]
                    df_filtered['Waktu'] = pd.to_datetime(df_filtered[kolom_timestamp], format='mixed', dayfirst=True).dt.strftime('%H:%M:%S')

                    kolom_inout = 'CLOCK IN / CLOCLK OUT' 

                    df_in = df_filtered[df_filtered[kolom_inout].str.contains('IN', na=False, case=False)].groupby('TANGGAL_ASLI')['Waktu'].min().reset_index()
                    df_in.rename(columns={'Waktu': 'In'}, inplace=True)

                    df_out = df_filtered[df_filtered[kolom_inout].str.contains('OUT', na=False, case=False)].groupby('TANGGAL_ASLI')['Waktu'].max().reset_index()
                    df_out.rename(columns={'Waktu': 'Out'}, inplace=True)

                    # Ekstraksi Shift
                    kolom_shift_gsheet = None
                    for col in df_filtered.columns:
                        if str(col).strip().upper() == 'SHIFT':
                            kolom_shift_gsheet = col
                            break

                    if kolom_shift_gsheet:
                        df_shift = df_filtered.dropna(subset=[kolom_shift_gsheet]).groupby('TANGGAL_ASLI')[kolom_shift_gsheet].max().reset_index()
                        df_shift['Shift'] = df_shift[kolom_shift_gsheet].apply(parse_shift)
                    else:
                        df_shift = pd.DataFrame({'TANGGAL_ASLI': df_filtered['TANGGAL_ASLI'].unique(), 'Shift': 'S1'})

                    df_rekap = pd.merge(df_in, df_out, on='TANGGAL_ASLI', how='outer')
                    df_rekap = pd.merge(df_rekap, df_shift[['TANGGAL_ASLI', 'Shift']], on='TANGGAL_ASLI', how='left')
                    df_rekap['Shift'] = df_rekap['Shift'].fillna('S1')
                    df_rekap = df_rekap.sort_values('TANGGAL_ASLI').reset_index(drop=True)

                    df_rekap['Date'] = df_rekap['TANGGAL_ASLI'].dt.strftime('%d-%b-%Y')
                    df_rekap['Day'] = df_rekap['TANGGAL_ASLI'].dt.day_name()
                    df_rekap['Type'] = "Working Day" 
                    df_rekap['BREAK'] = df_rekap['Day'].apply(lambda x: "1 Jam" if x not in ['Saturday', 'Sunday'] else "")

                    formulas = []
                    for i in range(len(df_rekap)):
                        row_excel = i + 2 
                        formula = f'=IF(OR(E{row_excel}="", G{row_excel}=""), "", ROUND(((G{row_excel}-E{row_excel})-TIME(1,0,0))*24, 1))'
                        formulas.append(formula)
                    
                    df_rekap['Working Hour(s)'] = formulas
                    df_rekap['Discription'] = ""

                    df_final = df_rekap[['Date', 'Day', 'Type', 'Shift', 'In', 'BREAK', 'Out', 'Working Hour(s)', 'Discription']]

                    # 3. GENERATE EXCEL KE DALAM MEMORY (TANPA SAVE KE HARDDISK SERVER)
                    buffer = io.BytesIO()
                    df_final.to_excel(buffer, index=False, engine='openpyxl')
                    buffer.seek(0)

                    # Formatting Openpyxl
                    wb = openpyxl.load_workbook(buffer)
                    ws = wb.active
                    center_alignment = Alignment(horizontal='center', vertical='center')

                    for col in ws.columns:
                        max_len = 0
                        col_letter = get_column_letter(col[0].column)
                        
                        for cell in col:
                            cell.alignment = center_alignment
                            if cell.value is not None:
                                val_str = str(cell.value)
                                if not val_str.startswith('='):
                                    max_len = max(max_len, len(val_str))
                                else:
                                    max_len = max(max_len, 16)
                        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

                    # Simpan hasil akhir format ke memory final
                    final_buffer = io.BytesIO()
                    wb.save(final_buffer)
                    final_buffer.seek(0)
                    
                    nama_file_excel = f"Rekap_Absen_{nama_input.replace(' ', '_')}_{tahun_input}_{bulan_input}.xlsx"

                    st.success(f"✅ Data {nama_input} berhasil diproses!")
                    
                    # 4. MUNCULKAN TOMBOL DOWNLOAD
                    st.download_button(
                        label="📥 Download File Excel",
                        data=final_buffer,
                        file_name=nama_file_excel,
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        type="primary"
                    )

            except KeyError as e:
                st.error(f"❌ Gagal! Cek nama kolom di Sheets: Kolom {e} tidak ditemukan.")
            except Exception as e:
                st.error(f"❌ Terjadi kesalahan: {e}")